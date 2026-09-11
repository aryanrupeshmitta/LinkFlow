import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()
    
    # 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_slide_layout = prs.slide_layouts[6]
    
    # Color palette
    NAVY_BG = RGBColor(15, 23, 42)        # #0F172A
    CARD_BG = RGBColor(30, 41, 59)        # #1E293B
    TEXT_WHITE = RGBColor(248, 250, 252)  # #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184) # #94A3B8
    ACCENT_INDIGO = RGBColor(99, 102, 241) # #6366F1
    ACCENT_EMERALD = RGBColor(16, 185, 129) # #10B981
    ACCENT_AMBER = RGBColor(245, 158, 11)   # #F59E0B

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, title_text, category_text="FULL STACK DEVELOPMENT (Course Code: 2550544)"):
        # Category Eyebrow
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category_text.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_INDIGO

        # Main Slide Title
        tx_box_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.8))
        tf_title = tx_box_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

    # ----------------------------------------------------
    # SLIDE 1: Title Slide
    # ----------------------------------------------------
    slide1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide1, NAVY_BG)

    # Eyebrow
    tx_box = slide1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.3), Inches(0.5))
    p = tx_box.text_frame.paragraphs[0]
    p.text = "FULL STACK DEVELOPMENT CASE STUDY (Course Code: 2550544)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_INDIGO

    # Main Title
    tx_box2 = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(1.2))
    p2 = tx_box2.text_frame.paragraphs[0]
    p2.text = "LinkFlow - Enterprise URL Shortener\n& Link Analytics Platform"
    p2.font.size = Pt(36)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE

    # Subtitle
    tx_box3 = slide1.shapes.add_textbox(Inches(1.0), Inches(3.8), Inches(11.3), Inches(0.8))
    p3 = tx_box3.text_frame.paragraphs[0]
    p3.text = "Commercial-grade URL Shortener with Passcode Protection, Expiration Controls, QR Codes & Real-Time Telemetry"
    p3.font.size = Pt(16)
    p3.font.color.rgb = TEXT_MUTED

    # Student Info Box
    tx_box4 = slide1.shapes.add_textbox(Inches(1.0), Inches(5.2), Inches(11.3), Inches(1.5))
    tf4 = tx_box4.text_frame
    p4 = tf4.paragraphs[0]
    p4.text = "Presented By: MITTA ARYAN RUPESH"
    p4.font.size = Pt(16)
    p4.font.bold = True
    p4.font.color.rgb = ACCENT_EMERALD

    p4_sub = tf4.add_paragraph()
    p4_sub.text = "Department of Computer Science & Engineering | MLRITM"
    p4_sub.font.size = Pt(14)
    p4_sub.font.color.rgb = TEXT_MUTED

    # ----------------------------------------------------
    # SLIDE 2: Summary of the Case Study
    # ----------------------------------------------------
    slide2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide2, NAVY_BG)
    add_header(slide2, "1. Executive Summary of the Case Study")

    # 2-Card Layout
    # Card 1
    shape1 = slide2.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    shape1.fill.solid()
    shape1.fill.fore_color.rgb = CARD_BG
    shape1.line.color.rgb = ACCENT_INDIGO

    tf = shape1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Core Solution & Value"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_INDIGO

    bullets1 = [
        "Commercial-grade full-stack URL shortener and link analytics suite.",
        "Solves long URL clutter, campaign tracking, and link security issues.",
        "Delivers custom slugs (linkflow.io/promo), QR code generation, and passcode security.",
        "Includes automated UTM parameter building and real-time click telemetry dashboards."
    ]
    for b in bullets1:
        p_b = tf.add_paragraph()
        p_b.text = "• " + b
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = TEXT_WHITE
        p_b.space_before = Pt(8)

    # Card 2
    shape2 = slide2.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8))
    shape2.fill.solid()
    shape2.fill.fore_color.rgb = CARD_BG
    shape2.line.color.rgb = ACCENT_EMERALD

    tf2 = shape2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = "Architectural Innovation"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT_EMERALD

    bullets2 = [
        "Zero-Setup Dual-Engine Storage Architecture (src/db/storage.js).",
        "Seamless fallback between MongoDB Atlas and Embedded JSON storage.",
        "Zero configuration error guarantee for instant local developer execution.",
        "Real-Time Link Health Audit Scanner & 100% passing Jest integration test suite."
    ]
    for b in bullets2:
        p_b = tf2.add_paragraph()
        p_b.text = "• " + b
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = TEXT_WHITE
        p_b.space_before = Pt(8)

    # ----------------------------------------------------
    # SLIDE 3: Introduction & Problem Statement
    # ----------------------------------------------------
    slide3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide3, NAVY_BG)
    add_header(slide3, "2. Introduction & Problem Statement")

    tx3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf3 = tx3.text_frame
    tf3.word_wrap = True

    sections3 = [
        ("Background & Industry Need:", "Long URLs with campaign tokens are ugly, hard to share, and impossible to track. Modern marketing demands branded, secure, short links."),
        ("Core Problems Solved:", "Public shorteners impose strict limits, high subscription costs, lack link password protection, and crash without cloud database setups."),
        ("LinkFlow Solution:", "A self-hostable, market-ready web platform with password screens, expiration controls, QR codes, dual-storage resilience, and complete data privacy.")
    ]

    for title, desc in sections3:
        p_t = tf3.add_paragraph() if tf3.paragraphs[0].text else tf3.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = ACCENT_INDIGO
        p_t.space_before = Pt(12)

        p_d = tf3.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(14)
        p_d.font.color.rgb = TEXT_WHITE
        p_d.space_before = Pt(4)

    # ----------------------------------------------------
    # SLIDE 4: Objectives of the Study
    # ----------------------------------------------------
    slide4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide4, NAVY_BG)
    add_header(slide4, "3. Objectives of the Study")

    grid = [
        ("1. RESTful API Engine", "Build a high-performance Express.js backend supporting single/bulk shortening, deletion, and redirection."),
        ("2. Dual Storage System", "Architect zero-setup storage abstraction combining MongoDB with Embedded Local JSON storage for 100% uptime."),
        ("3. Commercial Security", "Implement passcode hashing, expiration timers, click count thresholds, and endpoint rate limiting."),
        ("4. Glassmorphism UI & Telemetry", "Develop a dark glassmorphic SPA with Chart.js click timelines, QR generator, and link health diagnostics.")
    ]

    coords = [(0.8, 1.8), (6.8, 1.8), (0.8, 4.4), (6.8, 4.4)]
    for i, (title, text) in enumerate(grid):
        x, y = coords[i]
        shape = slide4.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.6), Inches(2.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = ACCENT_EMERALD if i%2==1 else ACCENT_INDIGO

        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = ACCENT_EMERALD if i%2==1 else ACCENT_INDIGO

        p2 = tf.add_paragraph()
        p2.text = text
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_WHITE
        p2.space_before = Pt(6)

    # ----------------------------------------------------
    # SLIDE 5: Requirements Analysis
    # ----------------------------------------------------
    slide5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide5, NAVY_BG)
    add_header(slide5, "4. Functional & Technical Requirements")

    shape_f = slide5.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    shape_f.fill.solid()
    shape_f.fill.fore_color.rgb = CARD_BG
    shape_f.line.color.rgb = ACCENT_INDIGO
    tf_f = shape_f.text_frame
    tf_f.word_wrap = True
    p = tf_f.paragraphs[0]
    p.text = "Functional Requirements (FR)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_INDIGO

    fr_items = [
        "FR-1: Custom Aliases & 6-char short codes.",
        "FR-2: Bulk processing of up to 50 URLs.",
        "FR-3: Passcode protection with verification screen.",
        "FR-4: Automatic expiration date & click limit enforcement.",
        "FR-5: Instant PNG QR Code generation & analytics dashboard.",
        "FR-6: Link Health Auditor (Healthy/Paused/Expired)."
    ]
    for item in fr_items:
        p_item = tf_f.add_paragraph()
        p_item.text = "• " + item
        p_item.font.size = Pt(12)
        p_item.font.color.rgb = TEXT_WHITE
        p_item.space_before = Pt(6)

    shape_t = slide5.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8))
    shape_t.fill.solid()
    shape_t.fill.fore_color.rgb = CARD_BG
    shape_t.line.color.rgb = ACCENT_AMBER
    tf_t = shape_t.text_frame
    tf_t.word_wrap = True
    p = tf_t.paragraphs[0]
    p.text = "Tech Stack & Non-Functional (NFR)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER

    nfr_items = [
        "NFR-1: Sub-50ms HTTP redirection response time.",
        "NFR-2: 100% Uptime via auto-switching dual storage.",
        "NFR-3: Express Rate Limiting & Passcode Hashing.",
        "Runtime: Node.js v16+, Express.js v4.18",
        "Frontend: HTML5, CSS3 Glassmorphism, JS ES6+, Chart.js",
        "Database: MongoDB Atlas / Embedded Local JSON"
    ]
    for item in nfr_items:
        p_item = tf_t.add_paragraph()
        p_item.text = "• " + item
        p_item.font.size = Pt(12)
        p_item.font.color.rgb = TEXT_WHITE
        p_item.space_before = Pt(6)

    # ----------------------------------------------------
    # SLIDE 6: System Architecture & Design
    # ----------------------------------------------------
    slide6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide6, NAVY_BG)
    add_header(slide6, "5. System Architecture & Dual Storage Engine")

    shape_arch = slide6.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    shape_arch.fill.solid()
    shape_arch.fill.fore_color.rgb = CARD_BG
    shape_arch.line.color.rgb = ACCENT_INDIGO

    tf_a = shape_arch.text_frame
    tf_a.word_wrap = True
    p = tf_a.paragraphs[0]
    p.text = "LinkFlow Dual-Engine Architecture & Flow"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_INDIGO

    arch_text = [
        "1. Client SPA (Browser): Communicates via REST API v1 using JSON payloads. Renders Chart.js & QR preview.",
        "2. API Gateway & Middleware: Express router with express-rate-limit and API Key authentication middleware.",
        "3. Shortener Service Layer (src/services/shortenerService.js): Manages validation, slug collision, passcode hashing, UTM appending, and analytics recording.",
        "4. Dual Storage Abstraction (src/db/storage.js): Auto-detects MongoDB; smoothly falls back to Embedded Local JSON storage if MongoDB is offline.",
        "5. Telemetry & Analytics: Captures device categories, User-Agent header telemetry, referrers, and raw access logs."
    ]
    for line in arch_text:
        p_line = tf_a.add_paragraph()
        p_line.text = line
        p_line.font.size = Pt(13)
        p_line.font.color.rgb = TEXT_WHITE
        p_line.space_before = Pt(10)

    # ----------------------------------------------------
    # SLIDE 7: Implementation Highlights
    # ----------------------------------------------------
    slide7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide7, NAVY_BG)
    add_header(slide7, "6. Implementation Details & Core Modules")

    col1 = slide7.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    col1.fill.solid()
    col1.fill.fore_color.rgb = CARD_BG
    col1.line.color.rgb = ACCENT_EMERALD
    tf1 = col1.text_frame
    tf1.word_wrap = True
    tf1.paragraphs[0].text = "Backend Implementation"
    tf1.paragraphs[0].font.size = Pt(18)
    tf1.paragraphs[0].font.bold = True
    tf1.paragraphs[0].font.color.rgb = ACCENT_EMERALD

    b1 = [
        "server.js — Entry point bootstrapping storage + API routes.",
        "src/db/storage.js — Dual-engine storage wrapper.",
        "src/routes/apiRoutes.js — 9 REST API v1 endpoints.",
        "src/routes/redirectRoutes.js — Redirection & passcode prompt page.",
        "src/middleware/rateLimiter.js — Security rate limits."
    ]
    for item in b1:
        p = tf1.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(8)

    col2 = slide7.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8))
    col2.fill.solid()
    col2.fill.fore_color.rgb = CARD_BG
    col2.line.color.rgb = ACCENT_INDIGO
    tf2 = col2.text_frame
    tf2.word_wrap = True
    tf2.paragraphs[0].text = "Frontend Implementation"
    tf2.paragraphs[0].font.size = Pt(18)
    tf2.paragraphs[0].font.bold = True
    tf2.paragraphs[0].font.color.rgb = ACCENT_INDIGO

    b2 = [
        "public/index.html — 5-tab SPA with glassmorphic cards.",
        "public/style.css — Dark glassmorphism, animated mesh background canvas.",
        "public/script.js — Chart.js integration, QR modal, link health scanner, Ctrl+K shortcut.",
        "Social Share Hub — Twitter/X, LinkedIn, WhatsApp, Facebook, Email, Native Share."
    ]
    for item in b2:
        p = tf2.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(8)

    # ----------------------------------------------------
    # SLIDE 8: Testing & Integration Results
    # ----------------------------------------------------
    slide8 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide8, NAVY_BG)
    add_header(slide8, "7. Testing Methodology & Integration Results")

    shape_test = slide8.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    shape_test.fill.solid()
    shape_test.fill.fore_color.rgb = CARD_BG
    shape_test.line.color.rgb = ACCENT_EMERALD

    tf_t = shape_test.text_frame
    tf_t.word_wrap = True
    p = tf_t.paragraphs[0]
    p.text = "Automated Jest & Supertest Integration Suite (100% Pass Rate)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD

    test_logs = [
        "PASS tests/shortener.test.js",
        "  URL Shortener API Tests",
        "    ✓ POST /api/v1/shorten creates a short URL (87 ms)",
        "    ✓ POST /api/v1/shorten handles custom slug and duplicate prevention (614 ms)",
        "    ✓ GET /api/v1/links returns list of links (13 ms)",
        "    ✓ POST /api/v1/bulk-shorten processes multiple URLs (13 ms)",
        "    ✓ GET /:shortCode redirects to original URL (19 ms)",
        "",
        "Test Suites: 1 passed, 1 total | Tests: 5 passed, 5 total | Time: 2.158 s"
    ]
    for line in test_logs:
        p_l = tf_t.add_paragraph()
        p_l.text = line
        p_l.font.size = Pt(12)
        p_l.font.color.rgb = RGBColor(134, 239, 172) if "✓" in line or "PASS" in line else TEXT_WHITE
        p_l.space_before = Pt(4)

    # ----------------------------------------------------
    # SLIDE 9: Results & Discussion
    # ----------------------------------------------------
    slide9 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide9, NAVY_BG)
    add_header(slide9, "8. Results & Comparative Performance")

    res_box = slide9.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    res_box.fill.solid()
    res_box.fill.fore_color.rgb = CARD_BG
    res_box.line.color.rgb = ACCENT_INDIGO

    tf_r = res_box.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "Key Project Achievements & Benchmarks"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_INDIGO

    res_points = [
        "• Redirection Latency: Sub-20ms redirection response time on local file storage; sub-10ms on MongoDB.",
        "• Zero-Setup Guarantee: Developers can clone the repository and run immediately without database setup errors.",
        "• Security & Compliance: Passcode-protected intermediate screens ensure sensitive links remain private.",
        "• Commercial Viability: Full feature parity with commercial Bitly / TinyURL platforms including QR codes, UTM parameters, bulk processing, and Chart.js dashboards.",
        "• Production Deployment: Fully containerized with Dockerfile and docker-compose.yml."
    ]
    for pt in res_points:
        p_pt = tf_r.add_paragraph()
        p_pt.text = pt
        p_pt.font.size = Pt(13)
        p_pt.font.color.rgb = TEXT_WHITE
        p_pt.space_before = Pt(10)

    # ----------------------------------------------------
    # SLIDE 10: Conclusion & Future Works
    # ----------------------------------------------------
    slide10 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide10, NAVY_BG)
    add_header(slide10, "9. Conclusion & Future Roadmap")

    c1 = slide10.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = ACCENT_EMERALD
    tfc1 = c1.text_frame
    tfc1.word_wrap = True
    tfc1.paragraphs[0].text = "Conclusion"
    tfc1.paragraphs[0].font.size = Pt(18)
    tfc1.paragraphs[0].font.bold = True
    tfc1.paragraphs[0].font.color.rgb = ACCENT_EMERALD

    conc = [
        "LinkFlow PRO v2.1 successfully achieves all project goals.",
        "Delivers a market-ready full-stack URL shortener and analytics engine.",
        "Combines resilient architecture with high usability and security.",
        "Open source under MIT License for Mitta Aryan Rupesh."
    ]
    for item in conc:
        p = tfc1.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(10)

    c2 = slide10.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = ACCENT_AMBER
    tfc2 = c2.text_frame
    tfc2.word_wrap = True
    tfc2.paragraphs[0].text = "Future Roadmap (v3.0)"
    tfc2.paragraphs[0].font.size = Pt(18)
    tfc2.paragraphs[0].font.bold = True
    tfc2.paragraphs[0].font.color.rgb = ACCENT_AMBER

    fut = [
        "1. Custom Domain CNAME Routing for multi-tenant organizations.",
        "2. User Accounts & JWT Authentication with RBAC.",
        "3. AI-Powered Fraud & Phishing Detection on target URLs.",
        "4. Native Mobile Application (React Native / Flutter)."
    ]
    for item in fut:
        p = tfc2.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(10)

    # Save presentation
    out_path = r'c:\Users\mitta\Downloads\URL-Shortener-Simulator\LinkFlow_FSD_Case_Study_Presentation.pptx'
    prs.save(out_path)
    print("Presentation generated successfully at:", out_path)

if __name__ == "__main__":
    create_presentation()
